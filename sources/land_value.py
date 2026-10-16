"""Written by AI on 2026.10.16 by model gpt-6.1-sol.

National land values excluding buildings, in 2024 market-exchange-rate US dollars.
OECD land/GDP ratios scale published years to 2024. Non-reporters use Fable's
fixed big-city-share model, fitted to the same fourteen whole-economy reporters.
Inputs and their download URLs are in data/raw/land-value/.
"""

from collections import Counter, defaultdict
import csv
import json
import math

import numpy as np
from openpyxl import load_workbook

from layers.grid import ROOT

RAW = ROOT / "data/raw/land-value"
PARTIAL_REPORTERS = {"MEX", "HRV", "NZL"}
HOUSEHOLD_SHARE = 0.6
NTD_PER_USD_2024 = 32.108

# UN M49 identifiers where the cached WB table has no 2024 GDP.
UN_GDP_CODES = {
    "AIA": 660, "CUB": 192, "ERI": 232, "GRL": 304, "KSV": 412,
    "MSR": 500, "PRK": 408, "SMR": 674, "SSD": 728, "SYR": 760,
    "VGB": 92, "YEM": 887,
}

# Conversion references only, without making a claim about sovereignty.
GDP_CONVERSION_REFERENCES = {
    "BES": "NLD", "BLM": "FRA", "FLK": "GBR", "GGY": "GBR",
    "GIB": "GBR", "GLP": "FRA", "GUF": "FRA", "JEY": "GBR",
    "MTQ": "FRA", "MYT": "FRA", "REU": "FRA", "SHN": "GBR",
    "SJM": "NOR", "SPM": "FRA", "TKL": "NZL", "WLF": "FRA",
    "ESH": "MAR", "ZNC": "TUR",
}


def world_bank_series(filename):
    series = defaultdict(dict)
    for record in json.loads((RAW / filename).read_text())[1]:
        if record["value"] is not None:
            series[record["countryiso3code"]][int(record["date"])] = record["value"]
    return series


def latest(series, year=None):
    if year in series:
        return year, series[year]
    years = [candidate for candidate in series if year is None or candidate <= year]
    return (max(years), series[max(years)]) if years else (None, None)


def official_land_ratios():
    local_gdp = world_bank_series("gdp_lcu.json")
    land = defaultdict(dict)
    with (RAW / "land_all.csv").open(newline="") as source:
        for row in csv.DictReader(source):
            if row["INSTR_ASSET"] != "N211N" or not row["OBS_VALUE"]:
                continue
            value = float(row["OBS_VALUE"]) * 10 ** int(row["UNIT_MULT"])
            if value > 0:
                land[row["REF_AREA"], row["SECTOR"]][int(row["TIME_PERIOD"])] = value
    ratios = {}
    for (country, sector), values in land.items():
        years = values.keys() & local_gdp.get(country, {}).keys()
        if years and sector in ("S1", "S1M") and country not in PARTIAL_REPORTERS:
            year = max(years)
            ratios[country, sector] = (year, values[year] / local_gdp[country][year])
    return ratios


def model_coefficients(ratios, urban_shares):
    countries = sorted(country for country, sector in ratios if sector == "S1")
    design = [[1.0, (latest(urban_shares.get(country, {}))[1] or 0.0) / 100]
              for country in countries]
    target = [math.log(ratios[country, "S1"][1]) for country in countries]
    return np.linalg.lstsq(design, target, rcond=None)[0].tolist()


def un_gdp():
    workbook = load_workbook(RAW / "un-gdp-usd.xlsx", read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    next(rows)
    next(rows)
    years = next(rows)
    codes = {number: code for code, number in UN_GDP_CODES.items()}
    result = {}
    for row in rows:
        if row[0] in codes and row[2] == "Gross Domestic Product (GDP)":
            series = {year: value for year, value in zip(years[3:], row[3:])
                      if isinstance(year, int) and isinstance(value, (float, int))}
            result[codes[row[0]]] = latest(series, 2024)
    workbook.close()
    return result


def taiwan_land_value():
    workbook = load_workbook(RAW / "dgbas-table1-2024.xlsx", read_only=True, data_only=True)
    for row in workbook.active.iter_rows(values_only=True):
        if row[0] == "End of 2024":
            total = row[11] * 1e12 / NTD_PER_USD_2024
            workbook.close()
            return total
    workbook.close()
    raise ValueError("DGBAS national wealth table has no 2024 land value")


def national_totals(country_codes):
    """Return a source-labelled 2024 land target for every GDP country code."""
    ratios = official_land_ratios()
    market_gdp = world_bank_series("gdp_usd.json")
    urban_shares = world_bank_series("wb_EN.URB.MCTY.TL.ZS.json")
    intercept, slope = model_coefficients(ratios, urban_shares)
    un_values = un_gdp()
    country_gdp = json.loads((ROOT / "data/build/gdp/conservation.json").read_text())["countries"]
    totals = {}
    for country in sorted(country_codes):
        record = {"total_usd2024": None, "source": None, "source_year": None,
                  "gdp_usd": None, "gdp_year": None, "gdp_source": None,
                  "big_city_share": None, "urban_share_missing": False,
                  "gdp_conversion_reference": None}
        if country == "TWN":
            record.update(total_usd2024=taiwan_land_value(), source="dgbas", source_year=2024)
            totals[country] = record
            continue
        gdp_year, gdp = latest(market_gdp.get(country, {}), 2024)
        gdp_source = "World Bank, GDP at current US dollars"
        if gdp_year != 2024 and country in un_values:
            gdp_year, gdp = un_values[country]
            gdp_source = "UN Statistics Division, GDP at current US dollars"
        if gdp is None:
            reference = GDP_CONVERSION_REFERENCES[country]
            gdp_year, reference_gdp = latest(market_gdp[reference], 2024)
            gdp = country_gdp[country]["source"] * reference_gdp / country_gdp[reference]["source"]
            gdp_source = "Kummu GDP scaled by reference country's market/PPP GDP ratio"
            record["gdp_conversion_reference"] = reference
        record.update(gdp_usd=gdp, gdp_year=gdp_year, gdp_source=gdp_source)
        if (country, "S1") in ratios:
            source_year, ratio = ratios[country, "S1"]
            source = "oecd_all_sectors"
        elif (country, "S1M") in ratios:
            source_year, ratio = ratios[country, "S1M"]
            ratio /= HOUSEHOLD_SHARE
            source = "oecd_household_scaled"
        else:
            source_year, city_percentage = latest(urban_shares.get(country, {}))
            record["urban_share_missing"] = city_percentage is None
            record["big_city_share"] = (city_percentage or 0.0) / 100
            ratio = math.exp(intercept + slope * record["big_city_share"])
            source = "big_city_model"
        record.update(total_usd2024=gdp * ratio, source=source, source_year=source_year)
        totals[country] = record
    return totals


def national_summary(records):
    ratios = official_land_ratios()
    urban_shares = world_bank_series("wb_EN.URB.MCTY.TL.ZS.json")
    country_gdp = json.loads((ROOT / "data/build/gdp/conservation.json").read_text())["countries"]
    proxy_countries = [country for country, record in records.items()
                       if record["gdp_conversion_reference"]]
    source_totals = defaultdict(float)
    for record in records.values():
        source_totals[record["source"]] += record["total_usd2024"]
    return {"country_count": len(records), "source_counts": dict(Counter(r["source"] for r in records.values())),
            "source_totals_usd2024": dict(source_totals),
            "world_total_usd2024": sum(source_totals.values()),
            "model_coefficients": model_coefficients(ratios, urban_shares),
            "household_share": HOUSEHOLD_SHARE,
            "missing_urban_share_assumption": "zero population in cities over one million",
            "missing_urban_share_countries": [c for c, r in records.items() if r["urban_share_missing"]],
            "prior_year_gdp_countries": {c: r["gdp_year"] for c, r in records.items()
                                         if r["gdp_year"] is not None and r["gdp_year"] < 2024},
            "gdp_conversion_proxy_countries": proxy_countries,
            "gdp_conversion_proxy_share_of_world_ppp_gdp": sum(country_gdp[c]["source"] for c in proxy_countries)
                / sum(r["source"] for r in country_gdp.values())}


def write_country_table(records, path):
    with path.open("w", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=["country", *next(iter(records.values())).keys()])
        writer.writeheader()
        writer.writerows({"country": country, **record} for country, record in records.items())
