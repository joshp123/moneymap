"""Written by AI on 2026.10.17 by model gpt-6.

Download and read the current GISCO and Census named-area boundaries.

    devenv shell -- python -m sources.named_areas
"""
from dataclasses import dataclass, field
import hashlib
import json
from pathlib import Path
import sqlite3
import urllib.error
import urllib.request
from zipfile import ZipFile

import fiona
from openpyxl import load_workbook
from rasterio.warp import transform_geom
import shapely
from shapely.geometry import shape

from layers.grid import ROOT

DIRECTORY = ROOT / 'data/sources/named-areas'
SOURCES = {
    'lau-2024.gpkg': 'https://gisco-services.ec.europa.eu/distribution/v2/lau/gpkg/LAU_RG_01M_2024_4326.gpkg',
    'nuts-2024.geojson': 'https://gisco-services.ec.europa.eu/distribution/v2/nuts/geojson/NUTS_RG_01M_2024_4326.geojson',
    'lau-nuts-2024.xlsx': 'https://ec.europa.eu/eurostat/documents/345175/501971/EU-27-LAU-2024-NUTS-2024.xlsx/12971f56-c035-dbab-4d9f-ff1dcc617bb3?t=1737036272863',
    **{f'us-{level}-2025.zip': f'https://www2.census.gov/geo/tiger/GENZ2025/shp/cb_2025_us_{level}_500k.zip'
       for level in ('place', 'county', 'state')},
}
EXCLUDED_US_STATES = {'60', '66', '69', '72', '78'}
# GISCO sovereign-country outlines include these separately owned grid countries.
GRID_OWNERS = {'FR': ('FRA', 'GLP', 'MTQ', 'GUF', 'REU', 'MYT', 'MAF'),
               'NO': ('NOR', 'SJM'), 'CY': ('CYP', 'ZNC')}
GRID_COUNTRY_ALIASES = {'XK': 'KSV'}


@dataclass
class Boundary:
    id: str
    code: str
    name: str
    level: str
    country: str
    source: str
    geometry: object
    parent_ids: list = field(default_factory=list)
    details: dict = field(default_factory=dict)


def validate(path):
    suffix = path.with_suffix('').suffix if path.suffix == '.download' else path.suffix
    if suffix in ('.zip', '.xlsx'):
        with ZipFile(path) as archive:
            if bad := archive.testzip():
                raise ValueError(f'Corrupt archive member: {bad}')
        if suffix == '.xlsx':
            workbook = load_workbook(path.open('rb'), read_only=True, data_only=True)
            count = len(workbook.sheetnames)
            workbook.close()
            return count
    elif suffix == '.gpkg':
        with sqlite3.connect(f'file:{path}?mode=ro', uri=True) as database:
            if database.execute('PRAGMA quick_check').fetchone()[0] != 'ok':
                raise ValueError(f'Corrupt GeoPackage: {path}')
    else:
        json.loads(path.read_text())
    with fiona.open('zip://' + str(path) if suffix == '.zip' else str(path)) as source:
        return len(source)


def download():
    DIRECTORY.mkdir(parents=True, exist_ok=True)
    records = []
    for name, url in SOURCES.items():
        path = DIRECTORY / name
        try:
            count = validate(path)
        except (OSError, ValueError, sqlite3.DatabaseError, fiona.errors.FionaError):
            print(f'Downloading {url}', flush=True)
            temporary = path.with_suffix(path.suffix + '.download')
            if path.exists() and not temporary.exists():
                path.replace(temporary)
            download_chunks(url, temporary)
            count = validate(temporary)
            temporary.replace(path)
        record = dict(file=name, url=url, bytes=path.stat().st_size,
                      sha256=hashlib.file_digest(path.open('rb'), 'sha256').hexdigest(),
                      **({'sheets': count} if path.suffix == '.xlsx' else {'features': count}))
        records.append(record)
        print(f'{name}: {count:,} {"sheets" if path.suffix == ".xlsx" else "features"}, {record["bytes"]:,} bytes', flush=True)
    return records


def download_chunks(url, destination):
    """Short resumable requests survive the source servers' long-stream resets."""
    chunk_size = 4 * 1024 * 1024
    total, failures = None, 0
    while total is None or not destination.exists() or destination.stat().st_size < total:
        offset = destination.stat().st_size if destination.exists() else 0
        request = urllib.request.Request(url, headers={'Range': f'bytes={offset}-{offset+chunk_size-1}'})
        try:
            with urllib.request.urlopen(request, timeout=120) as response:
                content_range = response.headers.get('Content-Range')
                if content_range:
                    total = int(content_range.rsplit('/', 1)[1])
                elif offset:
                    raise ValueError(f'Source does not honour resume requests: {url}')
                else:
                    total = int(response.headers['Content-Length'])
                with destination.open('ab') as output:
                    while block := response.read(1024 * 1024):
                        output.write(block)
                failures = 0
        except (OSError, urllib.error.URLError):
            failures += 1
            if failures == 3:
                raise


def polygon_geometry(geometry):
    geometry = shapely.make_valid(geometry) if not shapely.is_valid(geometry) else geometry
    if geometry.geom_type in ('Polygon', 'MultiPolygon'):
        return geometry
    polygons = [part for part in shapely.get_parts(geometry) if part.geom_type in ('Polygon', 'MultiPolygon')]
    return shapely.union_all(polygons)


def read_features(name):
    path = DIRECTORY / name
    with fiona.open('zip://' + str(path) if path.suffix == '.zip' else str(path)) as features:
        for feature in features:
            geometry = feature['geometry']
            if features.crs.to_epsg() != 4326:
                geometry = transform_geom(features.crs, 'EPSG:4326', geometry)
            yield dict(feature['properties']), polygon_geometry(shape(geometry))


def lau_nuts_correspondence():
    """Use Eurostat's validated administrative correspondence, including coastal LAUs."""
    workbook = load_workbook(DIRECTORY / 'lau-nuts-2024.xlsx', read_only=True, data_only=True)
    result = {}
    for sheet in workbook:
        if len(sheet.title) != 2:
            continue
        rows = sheet.iter_rows(values_only=True, max_col=3)
        header = next(rows)
        if header[0] not in ('NUTS3', 'NUTS 3 CODE') or header[1:] != ('LAU CODE', 'EU LAU CODE'):
            raise ValueError(f'Unexpected LAU correspondence columns: {sheet.title}: {header}')
        for nuts3, _, lau in rows:
            if nuts3 and lau:
                result.setdefault(str(lau).strip(), set()).add(str(nuts3).strip())
    workbook.close()
    return result


def load(countries=None):
    """Return boundaries, source-country metadata and hierarchy diagnostics."""
    selected = set(countries or ())
    boundaries, metadata, diagnostics = {}, {}, {'lau_parent_overlap': [], 'multicounty_places': [],
                                                 'lau_parent_methods': {'official_correspondence': 0, 'spatial_containment': 0}}
    nuts_rows = [] if selected == {'US'} else list(read_features('nuts-2024.geojson'))
    for properties, geometry in nuts_rows:
        country = properties['CNTR_CODE']
        if selected and country not in selected:
            continue
        code, level = properties['NUTS_ID'], int(properties['LEVL_CODE'])
        identifier = 'nuts-' + code
        source_iso3 = properties.get('ISO3_CODE') or properties.get('CNTR_ID')
        if not source_iso3:
            raise ValueError(f'GISCO has no ISO3 country for {code}')
        iso3 = GRID_COUNTRY_ALIASES.get(source_iso3, source_iso3)
        metadata[country] = dict(name=properties.get('NAME_ENGL') or country, iso3=iso3, source_iso3=source_iso3,
                                 grid_owners=list(GRID_OWNERS.get(country, (iso3,))))
        boundaries[identifier] = Boundary(identifier, code,
            properties.get('NAME_ENGL') if level == 0 else properties['NUTS_NAME'],
            'country' if level == 0 else f'nuts{level}', country, 'gisco', geometry,
            ['nuts-' + code[:length] for length in range(len(code)-1, 1, -1)])
    for country, record in metadata.items():
        if 'nuts-' + country not in boundaries:
            raise ValueError(f'GISCO missing country boundary: {country}')
    nuts3 = {country: [area for area in boundaries.values() if area.country == country and area.level == 'nuts3']
             for country in metadata}
    trees = {country: shapely.STRtree([area.geometry for area in areas]) for country, areas in nuts3.items() if areas}
    correspondence = {} if selected == {'US'} else lau_nuts_correspondence()
    for properties, geometry in (() if selected == {'US'} else read_features('lau-2024.gpkg')):
        country = properties['CNTR_CODE']
        if country not in metadata:
            continue
        code = properties['GISCO_ID']
        identifier = 'lau-' + code.replace('_', '-')
        parents = []
        official_codes = sorted(correspondence.get(code, ()))
        official_parents = [boundaries['nuts-' + value] for value in official_codes if 'nuts-' + value in boundaries]
        if official_parents and len(official_parents) == len(official_codes):
            parent = official_parents[0]
            parents = [candidate for candidate in [parent.id, *parent.parent_ids]
                       if all(candidate in [area.id, *area.parent_ids] for area in official_parents)]
            diagnostics['lau_parent_methods']['official_correspondence'] += 1
        elif country in trees:
            candidates = trees[country].query(geometry, predicate='intersects')
            if len(candidates):
                areas = nuts3[country]
                overlap = shapely.area(shapely.intersection([areas[i].geometry for i in candidates], geometry))
                parent = areas[candidates[int(overlap.argmax())]]
                fraction = float(max(overlap) / geometry.area) if geometry.area else 0
                chain = [parent.id, *parent.parent_ids]
                parents = [identifier for identifier in chain
                           if boundaries[identifier].level == 'country'
                           or shapely.intersection(boundaries[identifier].geometry, geometry).area / geometry.area >= .9999]
                skipped = [identifier for identifier in chain if identifier not in parents]
                if fraction < .99:
                    diagnostics['lau_parent_overlap'].append(dict(id=identifier, parent=parent.id, fraction=fraction))
                if skipped:
                    diagnostics.setdefault('lau_skipped_parents', []).append(dict(id=identifier, name=properties['LAU_NAME'],
                        skipped=skipped, parents=parents, largest_nuts3_overlap=fraction))
            diagnostics['lau_parent_methods']['spatial_containment'] += 1
        if not parents:
            parents = ['nuts-' + country]
        boundaries[identifier] = Boundary(identifier, code, properties['LAU_NAME'], 'lau', country,
                                           'gisco', geometry, parents)
    if not selected or 'US' in selected:
        metadata['US'] = dict(name='United States', iso3='USA', grid_owners=['USA'])
        us_states, counties = {}, []
        for properties, geometry in read_features('us-state-2025.zip'):
            code = properties['STATEFP']
            if code in EXCLUDED_US_STATES:
                continue
            identifier = 'us-state-' + code
            us_states[code] = Boundary(identifier, code, properties['NAME'], 'state', 'US', 'census', geometry, ['us-country'])
            boundaries[identifier] = us_states[code]
        boundaries['us-country'] = Boundary('us-country', 'US', 'United States', 'country', 'US', 'census',
                                             shapely.union_all([area.geometry for area in us_states.values()]))
        for properties, geometry in read_features('us-county-2025.zip'):
            state = properties['STATEFP']
            if state not in us_states:
                continue
            code = properties['GEOID']
            county = Boundary('us-county-' + code, code, properties.get('NAMELSAD') or properties['NAME'],
                              'county', 'US', 'census', geometry, ['us-state-' + state, 'us-country'])
            boundaries[county.id] = county
            counties.append(county)
        county_tree = shapely.STRtree([area.geometry for area in counties])
        for properties, geometry in read_features('us-place-2025.zip'):
            state = properties['STATEFP']
            if state not in us_states:
                continue
            code = properties['GEOID']
            identifier = 'us-place-' + code
            parents = ['us-state-' + state, 'us-country']
            candidates = county_tree.query(geometry, predicate='intersects')
            if len(candidates):
                overlap = shapely.area(shapely.intersection([counties[i].geometry for i in candidates], geometry))
                maximum = int(overlap.argmax())
                fraction = float(overlap[maximum] / geometry.area) if geometry.area else 0
                if fraction > .99999:
                    parents.insert(0, counties[candidates[maximum]].id)
                else:
                    diagnostics['multicounty_places'].append(dict(id=identifier, name=properties['NAME'],
                        counties=[dict(id=counties[i].id, fraction=float(share / geometry.area))
                                  for i, share in zip(candidates, overlap) if share / geometry.area > .00001]))
            boundaries[identifier] = Boundary(identifier, code, properties['NAME'], 'place', 'US', 'census',
                                               geometry, parents, dict(census_lsad=properties['LSAD']))
    for boundary in boundaries.values():
        boundary.parent_ids = [identifier for identifier in boundary.parent_ids if identifier in boundaries]
    return boundaries, metadata, diagnostics


if __name__ == '__main__':
    download()
